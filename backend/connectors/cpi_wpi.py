"""
CPI and WPI connectors -- macro inflation benchmarks used ONLY for
comparison against the observed household basket (frozen PRD section 7/17:
these must never be presented as the same thing as the observed basket).

CPI (MOSPI): no stable JSON API exists, monthly PDF releases only -- CSV_PATH
manual update remains the honest V1 approach here.

WPI (OEA): CORRECTED -- a real, live, no-login .xlsx download exists at
eaindustry.nic.in (verified Sept 2026, see WPIConnector.fetch() docstring),
so this is no longer a manual-CSV connector. It fetches automatically.
"""
import os
import csv
import requests
from .base import BaseConnector


class CPIConnector(BaseConnector):
    source_name = "CPI (MOSPI)"
    source_type = "government"
    quality_grade_default = "A"

    def fetch(self, **kwargs):
        path = os.environ.get("CPI_CSV_PATH")
        if not path or not os.path.exists(path):
            return None
        with open(path) as f:
            return list(csv.DictReader(f))

    def parse(self, raw):
        return raw or []

    def normalize(self, parsed_rows):
        return parsed_rows

    def validate(self, rows):
        """Override: base class's validate() checks normalized_unit_price,
        which price_observations rows have but inflation_benchmarks rows
        (CPI/WPI) never will -- that mismatch was silently rejecting every
        real CPI/WPI row regardless of CSV quality. Check what this
        connector's rows actually need instead: a numeric index value and a
        valid month/year."""
        good = []
        rejected = 0
        for r in rows:
            try:
                float(r["value"])
                int(r["month"])
                int(r["year"])
            except (KeyError, ValueError, TypeError):
                rejected += 1
                continue
            good.append(r)
        return good, rejected

    def save(self, rows):
        cur = self.conn.cursor()
        source_id = self.get_source_id()
        saved = 0
        for r in rows:
            try:
                cur.execute(
                    """INSERT INTO inflation_benchmarks
                       (benchmark_type, commodity_group, state, month, year, value, mom_change, yoy_change, status, source_id)
                       VALUES (?,?,?,?,?,?,?,?,?,?)""",
                    (r.get("benchmark_type", "CPI"), r.get("commodity_group", "Headline"), r.get("state"),
                     int(r["month"]), int(r["year"]), float(r["value"]),
                     float(r["mom_change"]) if r.get("mom_change") else None,
                     float(r["yoy_change"]) if r.get("yoy_change") else None,
                     r.get("status", "final"), source_id),
                )
                saved += 1
            except (KeyError, ValueError, TypeError):
                continue
        self.conn.commit()
        return saved


class WPIConnector(BaseConnector):
    source_name = "WPI (OEA)"
    source_type = "government"
    quality_grade_default = "A"

    DOWNLOAD_PAGE = "https://eaindustry.nic.in/download_data_2223.asp"

    def fetch(self, **kwargs):
        """
        REAL, verified live source (confirmed reachable Sept 2026, no login,
        no key): the Office of Economic Adviser publishes the current
        (2022-23 base) WPI monthly index as a plain .xlsx, no auth needed.
        The exact filename changes every month when a new release is
        published (e.g. wpi_monthly_index_202608.xlsx for Aug 2026), so we
        scrape the current link off the download page rather than guessing
        the filename -- that page itself is not robots-disallowed.

        Honest caveat: the download page's HTML structure and the .xlsx's
        internal row/column layout were confirmed via page content, not by
        opening the actual binary file (not possible from this environment).
        parse() below searches for a row labelled "All Commodities" rather
        than a hardcoded cell reference, specifically to tolerate that
        uncertainty -- but the first real run should be checked by hand
        against the actual downloaded file to confirm columns line up.
        """
        try:
            page = requests.get(self.DOWNLOAD_PAGE, timeout=30)
            page.raise_for_status()
        except requests.exceptions.Timeout:
            print("WPI: download page request timed out after 30s")
            return None
        except requests.exceptions.ConnectionError as e:
            print(f"WPI: connection failed reaching download page -- {e}")
            return None
        except requests.exceptions.HTTPError as e:
            print(f"WPI: download page returned HTTP error -- {e}")
            return None
        except Exception as e:
            print(f"WPI: unexpected error fetching download page -- {type(e).__name__}: {e}")
            return None

        import re
        match = re.search(r'href="(https://eaindustry\.nic\.in/indx_download_2223/wpi_monthly_index_\d+\.xlsx)"', page.text)
        if not match:
            print("WPI: download page fetched OK, but couldn't find the .xlsx link pattern in it -- page structure may have changed")
            return None
        xlsx_url = match.group(1)
        print(f"WPI: found xlsx link -- {xlsx_url}")

        try:
            resp = requests.get(xlsx_url, timeout=30)
            resp.raise_for_status()
        except requests.exceptions.Timeout:
            print("WPI: .xlsx download timed out after 30s")
            return None
        except requests.exceptions.ConnectionError as e:
            print(f"WPI: connection failed downloading .xlsx -- {e}")
            return None
        except requests.exceptions.HTTPError as e:
            print(f"WPI: .xlsx download returned HTTP error -- {e}")
            return None
        except Exception as e:
            print(f"WPI: unexpected error downloading xlsx -- {type(e).__name__}: {e}")
            return None

        import io
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(resp.content), read_only=True, data_only=True)
            print(f"WPI: xlsx downloaded and parsed successfully ({len(resp.content)} bytes)")
        except Exception as e:
            print(f"WPI: downloaded file but openpyxl couldn't parse it -- {type(e).__name__}: {e}")
            return None
        return wb

    def parse(self, raw):
        """Search for the 'All Commodities' headline row rather than assume
        a fixed cell position -- see fetch() docstring for why."""
        if raw is None:
            return []
        rows = []
        try:
            ws = raw.active
            header_row = None
            for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
                if row and any(isinstance(c, str) and c.strip().lower() in ("month", "month/year") for c in row):
                    header_row = row
                    break
            if header_row is None:
                return []
            for row in ws.iter_rows(values_only=True):
                if row and isinstance(row[0], str) and "all commodities" in row[0].strip().lower():
                    for col_idx, val in enumerate(row):
                        if col_idx == 0 or val is None:
                            continue
                        header = header_row[col_idx] if col_idx < len(header_row) else None
                        if header:
                            rows.append({"period_label": str(header), "value": val})
                    break
        except Exception:
            return []
        return rows

    def normalize(self, parsed_rows):
        """Best-effort period_label -> month/year; genuinely ambiguous
        formats are dropped rather than guessed, per the frozen PRD.
        Handles both 4-digit (2025) and 2-digit (25) year formats since the
        real file's exact convention wasn't verifiable from this
        environment -- confirm which one actually appears on first real run."""
        import re
        out = []
        for r in parsed_rows:
            label = str(r.get("period_label", ""))
            m = re.search(r"(\d{1,2})[/\-](\d{2,4})\b|([A-Za-z]{3,})[\s\-](\d{2,4})\b", label)
            if not m:
                continue
            try:
                if m.group(1):
                    month, year = int(m.group(1)), int(m.group(2))
                else:
                    from datetime import datetime
                    month = datetime.strptime(m.group(3)[:3], "%b").month
                    year = int(m.group(4))
                if year < 100:
                    year += 2000
            except (ValueError, TypeError):
                continue
            try:
                value = float(r["value"])
            except (TypeError, ValueError):
                continue
            out.append({"benchmark_type": "WPI", "commodity_group": "All Commodities",
                        "state": None, "month": month, "year": year, "value": value,
                        "mom_change": None, "yoy_change": None, "status": "final"})
        return out

    def validate(self, rows):
        """Same override as CPIConnector -- see that class for why the base
        validate() was silently rejecting every row here too."""
        good = []
        rejected = 0
        for r in rows:
            try:
                float(r["value"])
                int(r["month"])
                int(r["year"])
            except (KeyError, ValueError, TypeError):
                rejected += 1
                continue
            good.append(r)
        return good, rejected

    def save(self, rows):
        cur = self.conn.cursor()
        source_id = self.get_source_id()
        saved = 0
        for r in rows:
            try:
                cur.execute(
                    """INSERT INTO inflation_benchmarks
                       (benchmark_type, commodity_group, state, month, year, value, mom_change, yoy_change, status, source_id)
                       VALUES (?,?,?,?,?,?,?,?,?,?)""",
                    ("WPI", r.get("commodity_group", "Food"), r.get("state"),
                     int(r["month"]), int(r["year"]), float(r["value"]),
                     float(r["mom_change"]) if r.get("mom_change") else None,
                     float(r["yoy_change"]) if r.get("yoy_change") else None,
                     r.get("status", "final"), source_id),
                )
                saved += 1
            except (KeyError, ValueError, TypeError):
                continue
        self.conn.commit()
        return saved
